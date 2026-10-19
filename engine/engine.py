"""Scenario engine: a Python replica of the workbook's operating model and DCF, as a function of its inputs.
Terminal margin (15 Assumptions rows 93-106) and WACC (20 DCF Valuation H8-H18) are recomputed from
their own inputs, so any of them can be moved. load(path) reads inputs once; value(case) prices one case."""
import copy
import openpyxl

SEG_ROWS = (10, 11, 12, 13, 14, 15)          # model rows: Search, YouTube, Network, Subs, Cloud, Other Bets
DRV_ROW = {10: 18, 11: 19, 12: 20, 13: 21, 14: 22, 15: 23}


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    A, D, V = wb['15 Assumptions'], wb['20 DCF Valuation'], wb['11 Valuation']
    a = lambda c: A[c].value
    base = dict(
        t=a('C10'), tg=a('C85'), sbc=[a(f'{c}45') for c in 'CDEFG'], oth=[a(f'{c}46') for c in 'CDEFG'],
        gain=[a(f'{c}47') or 0 for c in 'CDEFG'], c24=a('C24') or 0, nis0=a('C36'), lag=a('C42'), c56=a('C56'),
        c58=a('C58'), c48=a('C48'), c57=a('C57') or 0, c59=a('C59'), shares0=a('C8'), price=a('C7'), pref_div=a('C9'),
        N=a('C79'), life_s=a('C40'), life_l=a('C41'),
        rf=D['H8'].value, erp=D['H9'].value, beta=D['H10'].value, kd=D['H13'].value, dw=D['H17'].value,
        g=D['H19'].value, H20=D['H20'].value, netcash=D['H21'].value, pref=D['H22'].value, dil=D['H23'].value,
        stakes=D['C25'].value, days=(V['C87'].value - D['H25'].value).days,
    )
    M = wb['DCF Base Model']
    base['seg0'] = {r: M.cell(r, 10).value for r in SEG_ROWS}
    base['base_rate'] = M.cell(64, 11).value / M.cell(23, 11).value
    cases, checks = {}, {}
    for name, dcol in (('Base', 'C'), ('Best', 'D'), ('Worst', 'E')):
        DR = wb[f'DCF {name} Drivers']
        dr = lambda r, i: DR.cell(r, 3 + i).value
        c = copy.deepcopy(base)
        c['growth'] = {r: [dr(DRV_ROW[r], i) for i in range(5)] for r in SEG_ROWS}
        c['cents'] = [None] + [dr(24, i) for i in range(1, 5)]
        c['margin'] = [dr(25, i) for i in range(5)]
        c['capex'] = [dr(26, i) for i in range(5)]
        c['ronic'] = DR['C27'].value
        c['short'] = DR['C28'].value
        c['name'] = name
        cases[name] = c
        checks[name] = (D[f'{dcol}29'].value, V[f'{dcol}90'].value)
    return cases, checks


def wacc_of(c):
    ke = c['rf'] + c['beta'] * c['erp']
    return (1 - c['dw']) * ke + c['dw'] * c['kd'] * (1 - c['t']), ke


def terminal_margin(c, m30_prerramp, life):
    g, r, t = c['g'], c['ronic'], c['t']
    dep_capex = (1 + g) ** -(c['lag'] - 0.5) * (1 - (1 + g) ** -life) / (g * life)
    ci = (g / r) * (1 - t) * (m30_prerramp + c['base_rate']) / ((1 - dep_capex) + (g / r) * (1 - t) * dep_capex)
    return m30_prerramp + c['base_rate'] - ci * dep_capex


def value(c, terminal='derived'):
    """Returns per-share value as built (Jun 30, end-of-period), at the price date (mid-period, rolled at ke),
    and the yearly paths. terminal: 'derived' (base method) or 'held' (FY2030E margin held)."""
    t = c['t']
    life = 1 / (c['short'] / c['life_s'] + (1 - c['short']) / c['life_l'])
    seg = dict(c['seg0'])
    res = []
    prev_capex = None; nis_prev = c['nis0']; cum = 0.0; shares = None
    for i in range(5):
        for r in (10, 11, 12, 13, 15):
            seg[r] *= 1 + c['growth'][r][i]
        seg[14] = seg[14] * (1 + c['growth'][14][0]) if i == 0 else seg[14] + prev_capex * c['cents'][i]
        rev = sum(seg.values()) + c['c24']
        capex = rev * c['capex'][i]
        nis = c['lag'] * capex
        if i == 0:
            placed = capex - c['c56'] - (nis - nis_prev); cum = placed
            charge = placed * c['c58'] / life
        else:
            placed = capex - (nis - nis_prev)
            charge = (cum + placed / 2) / life; cum += placed
        nis_prev = nis
        dep = c['c48'] + charge
        ebit = rev * c['margin'][i] - (dep - rev * c['base_rate'])
        oth, gain, sbc = c['oth'][i], c['gain'][i], rev * c['sbc'][i]
        core_pt = ebit + oth
        tax = core_pt * t + gain * c['tg']
        ni = ebit + oth + gain - tax
        shares = c['shares0'] + 0.25 * sbc / c['price'] if i == 0 else shares + sbc / c['price']
        eps_core = (core_pt * (1 - t) - c['pref_div']) / shares
        one_off = c['c59'] if i == 0 else 0
        ocf = ni + dep + sbc - gain + rev * c['c57'] + gain * c['tg'] + one_off
        fcff = ebit * (1 - t) + dep + rev * c['c57'] + one_off - capex
        res.append(dict(rev=rev, ebit=ebit, dep=dep, capex=capex, eps_core=eps_core, fcf=ocf - capex, fcff=fcff,
                        cloud=seg[14], margin=ebit / rev))
        prev_capex = capex
    wacc, ke = wacc_of(c)
    g, N = c['g'], c['N']
    g30 = res[4]['rev'] / res[3]['rev'] - 1
    m30 = res[4]['ebit'] / res[4]['rev']
    tm = terminal_margin(c, c['margin'][4], life) if terminal == 'derived' else m30
    rate30 = (res[4]['capex'] - res[4]['dep'] - res[4]['rev'] * c['c57']) / (res[4]['ebit'] * (1 - t))
    rev = res[4]['rev']; pv_fade = 0.0; last = None; fade = []
    for k in range(1, 11):
        gk = g30 + (g - g30) * k / N if k <= N else g
        rev *= 1 + gk
        nopat = rev * (m30 + (tm - m30) * min(k, N) / N) * (1 - t)
        rr = g / c['ronic'] if k >= N else rate30 + (g / c['ronic'] - rate30) * (k - 1) / max(1, N - 1)
        last = nopat * (1 - rr)
        fade.append(dict(rev=rev, fcff=last))
        pv_fade += last / (1 + wacc) ** (k + 4.5)
    pv = (res[0]['fcff'] - c['H20']) / (1 + wacc) ** 0.5 + sum(res[i]['fcff'] / (1 + wacc) ** (i + 0.5) for i in range(1, 5))
    tv = last * (1 + g) / (wacc - g) / (1 + wacc) ** 14.5
    bridge = c['netcash'] - c['pref'] + c['stakes']
    vps = (pv + pv_fade + tv + bridge) / c['dil']
    mid = ((res[0]['fcff'] - c['H20']) / (1 + wacc) ** 0.25 + sum(res[i]['fcff'] / (1 + wacc) ** i for i in range(1, 5))
           + pv_fade * (1 + wacc) ** 0.5 + tv + bridge) / c['dil']
    pdv = mid * (1 + ke) ** (c['days'] / 365)
    return dict(vps=vps, pdv=pdv, years=res, fade=fade, wacc=wacc, ke=ke, tmargin=tm, tv_share=tv / (pv + pv_fade + tv))


if __name__ == '__main__':
    import sys
    cases, checks = load(sys.argv[1])
    for n, c in cases.items():
        r = value(c)
        print(f"{n:5s} as built {r['vps']:.6f} vs {checks[n][0]:.6f} | price date {r['pdv']:.6f} vs {checks[n][1]:.6f} | "
              f"diff {max(abs(r['vps']/checks[n][0]-1), abs(r['pdv']/checks[n][1]-1)):.1e} | tm {r['tmargin']:.4f} wacc {r['wacc']:.5f}")
        print('      held-margin price-date', round(value(c, 'held')['pdv'], 2))
